"""
document_processor.py
----------------------
Extracts raw text from uploaded learning material: PDF, DOCX, PPTX, TXT.
This is step 1 of the RAG pipeline described in the assessment
("Learning Material Processing").
"""

import os
from PyPDF2 import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation


def extract_text(file_path: str) -> str:
    """Detects file type from extension and extracts plain text."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    elif ext == ".pptx":
        return _extract_pptx(file_path)
    elif ext in (".txt", ".md"):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    return "\n".join(pages)


def _extract_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # also pull text out of tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def _extract_pptx(file_path: str) -> str:
    prs = Presentation(file_path)
    slides_text = []
    for i, slide in enumerate(prs.slides, start=1):
        chunk = [f"[Slide {i}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        chunk.append(line)
        slides_text.append("\n".join(chunk))
    return "\n\n".join(slides_text)


def chunk_text(text: str, chunk_size: int = 900, overlap: int = 150):
    """Splits long text into overlapping chunks for retrieval."""
    text = " ".join(text.split())  # normalize whitespace
    if not text:
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start = end - overlap
    return chunks
